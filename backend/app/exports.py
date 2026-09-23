"""Local, Unicode DOCX/PDF exports of the approved proposal."""
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Pt
from fontTools.ttLib import TTFont
from fpdf import FPDF

from backend.shared.schemas import Proposal
from .config import Settings
from .models import Run


def pdf_font(settings: Settings) -> Path:
    candidates = [Path(settings.pdf_font_path)] if settings.pdf_font_path else []
    candidates += [Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                   Path("/usr/local/share/fonts/DejaVuSans.ttf"),
                   Path("/System/Library/Fonts/Supplemental/Georgia.ttf"),
                   Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
                   Path("C:/Windows/Fonts/DejaVuSans.ttf"), Path("C:/Windows/Fonts/arial.ttf")]
    required = set(map(ord, "әғқңөұүһіӘҒҚҢӨҰҮҺІ"))
    for path in candidates:
        if path.is_file():
            with TTFont(path) as font:
                if required.issubset(font.getBestCmap()):
                    return path
    raise RuntimeError("Не найден шрифт с казахскими буквами. Задайте PDF_FONT_PATH (TTF).")


def paragraphs(run: Run, proposal: Proposal):
    yield "Протокол совещания"
    yield run.title
    date_line = run.meeting_date.isoformat() if run.meeting_date else "не указана"
    yield f"Дата: {date_line}" + ("" if run.meeting_date_verified or not run.meeting_date else " (не подтверждена)")
    if run.synthetic:
        yield "Синтетические демонстрационные данные. Имена и содержание вымышлены."
    yield "Участники: " + ", ".join(p["name"] for p in run.participants)
    yield "Резюме"
    yield proposal.summary
    yield "Решения"
    for decision in proposal.decisions:
        yield decision
    yield "Поручения"
    for i, item in enumerate([a for a in proposal.assignments if a.review_status != "excluded"], 1):
        yield f"{i}. {item.assignee or 'Ответственный не указан'}: {item.task}"
        yield f"Срок: {item.deadline or 'не указан'}; приоритет: {item.priority}."
        if item.deadline_text:
            yield f"Формулировка срока: {item.deadline_text}"
    yield "Транскрипт"
    segments = proposal.segments or run.segments
    for raw in segments:
        segment = raw.model_dump() if hasattr(raw, "model_dump") else raw
        speaker = proposal.speakers.get(segment["speaker"] or "", segment["speaker"]) or "Говорящий не определён"
        text = segment.get("corrected_text") or segment["text"]
        yield f"[{segment['start']:.1f}–{segment['end']:.1f}] {speaker}: {text}"


def export_protocol(run: Run, settings: Settings, kind: str) -> Path:
    # Exports reflect only the approved snapshot; run.proposal remains for legacy/seed rows.
    proposal = Proposal.model_validate(run.approved or run.proposal)
    directory = settings.exports_dir / run.id
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / f"protocol.{kind}"
    temporary = directory / f"{uuid4().hex}.{kind}"
    lines = list(paragraphs(run, proposal))
    try:
        if kind == "docx":
            document = Document()
            document.styles["Normal"].font.name = "Arial"
            document.styles["Normal"].font.size = Pt(11)
            for i, line in enumerate(lines):
                if i == 0:
                    document.add_heading(line, 0)
                elif line in {"Резюме", "Решения", "Поручения", "Транскрипт"}:
                    document.add_heading(line, 1)
                else:
                    document.add_paragraph(line)
            document.save(str(temporary))
        elif kind == "pdf":
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_font("Protocol", fname=str(pdf_font(settings)))
            pdf.add_page()
            pdf.set_font("Protocol", size=11)
            for line in lines:
                pdf.multi_cell(0, 6, text=line, new_x="LMARGIN", new_y="NEXT")
                pdf.ln(2)
            pdf.output(str(temporary))
        else:
            raise ValueError("Неизвестный формат протокола.")
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return target
